"""Data-preservation, orientation and rejection tests for the legacy importer."""

from __future__ import annotations

import copy
import importlib.util
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest

import numpy as np
import openpyxl
import pandas as pd
from scipy import stats


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "analysis" / "legacy_workbook.py"
SPEC = importlib.util.spec_from_file_location("legacy_workbook_under_test", SOURCE)
legacy = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(legacy)


def source_rows():
    return [
        ["ENSMUSG00000000001", "Shared", 100, 100.0, 0.4, 0.1, 0.001, 0.003],
        ["ENSMUSG00000000002", "Shared", 100, 20.0, -0.5, 0.2, 0.01, 0.015],
        ["ENSMUSG00000000003", "Other", 200, 5.0, 0.0, 0.05, 0.5, 0.5],
    ]


def workbook_bytes(rows=None, *, headers=None, extra_sheet=False):
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for _, sheet_name, _, _, _ in legacy.FIGURES:
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(list(headers or legacy.HEADERS))
        for row in copy.deepcopy(rows if rows is not None else source_rows()):
            sheet.append(row)
    if extra_sheet:
        workbook.create_sheet("unexpected")
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class SourceValidationTests(unittest.TestCase):
    def test_all_three_source_sheets_are_loaded_without_collapsing_ensembl_rows(self):
        tables = legacy.read_workbook(workbook_bytes())
        self.assertEqual(set(tables), {item[1] for item in legacy.FIGURES})
        for table in tables.values():
            self.assertEqual(len(table), 3)
            self.assertEqual(table.SYMBOL.duplicated().sum(), 1)
            self.assertEqual(table.ENTREZID.duplicated().sum(), 1)
            self.assertEqual(table.ENSEMBL.nunique(), 3)

    def test_unexpected_or_missing_sheets_are_rejected(self):
        with self.assertRaisesRegex(ValueError, "Expected exactly"):
            legacy.read_workbook(workbook_bytes(extra_sheet=True))
        workbook = openpyxl.load_workbook(BytesIO(workbook_bytes()))
        workbook.remove(workbook["saline_VS_wt"])
        buffer = BytesIO()
        workbook.save(buffer)
        with self.assertRaisesRegex(ValueError, "Expected exactly"):
            legacy.read_workbook(buffer.getvalue())

    def test_header_changes_are_rejected(self):
        headers = list(legacy.HEADERS)
        headers[-1] = "qvalue"
        with self.assertRaisesRegex(ValueError, "eight documented column headers"):
            legacy.read_workbook(workbook_bytes(headers=headers))

    def test_missing_numeric_and_formula_cells_are_rejected(self):
        for value in (None, "NA", "=1/2", True):
            rows = source_rows()
            rows[0][6] = value
            with self.subTest(value=value), self.assertRaises(ValueError):
                legacy.read_workbook(workbook_bytes(rows))

    def test_probability_bounds_and_underflow_are_not_silently_corrected(self):
        for column, value in ((6, 0), (6, -0.1), (6, 1.1), (7, 0), (7, 1.1)):
            rows = source_rows()
            rows[0][column] = value
            with self.subTest(column=column, value=value), self.assertRaisesRegex(ValueError, "underflow"):
                legacy.read_workbook(workbook_bytes(rows))

    def test_invalid_counts_errors_annotations_and_duplicate_ids_are_rejected(self):
        for column, value in ((3, -1), (5, 0), (5, -0.1), (2, 100.5), (1, ""), (0, "ENSG000001")):
            rows = source_rows()
            rows[0][column] = value
            with self.subTest(column=column, value=value), self.assertRaises(ValueError):
                legacy.read_workbook(workbook_bytes(rows))
        rows = source_rows()
        rows[1][0] = rows[0][0]
        with self.assertRaisesRegex(ValueError, "duplicate Ensembl"):
            legacy.read_workbook(workbook_bytes(rows))


class OrientationAndImportTests(unittest.TestCase):
    def test_leptin_vs_saline_reverses_only_effect_direction_and_proxy_sign(self):
        source = legacy.read_workbook(workbook_bytes())["saline_VS_leptin"]
        oriented = legacy.orient_table(source, "saline_VS_leptin", -1)
        np.testing.assert_array_equal(oriented.log2FoldChange, -source.log2FoldChange)
        np.testing.assert_array_equal(oriented.source_log2FoldChange, source.log2FoldChange)
        for column in ("baseMean", "lfcSE", "pvalue", "padj"):
            np.testing.assert_array_equal(oriented[column], source[column])
        signs = -np.sign(source.log2FoldChange)
        np.testing.assert_allclose(oriented.primaryproxy_signed_log10_p, signs * -np.log10(source.pvalue))
        np.testing.assert_allclose(oriented.sensitivityproxy_signed_z, signs * stats.norm.isf(source.pvalue / 2))
        self.assertEqual(oriented.supplied_symbol.tolist(), ["Shared", "Shared", "Other"])

    def test_import_preserves_source_probabilities_and_records_confirmation_hashes(self):
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "source.xlsx"
            source.write_bytes(workbook_bytes())
            destination = Path(temporary) / "imported"
            manifest = legacy.import_workbook(source, destination, complete_tested_universe=True)
            self.assertTrue(manifest["complete"])
            self.assertTrue(manifest["tested_universe_confirmation"]["all_tested_genes_exported"])
            self.assertEqual(manifest["input_file"], "source.xlsx")
            self.assertNotIn(temporary, json.dumps(manifest))
            self.assertFalse(manifest["preservation"]["source_pvalues_recomputed"])
            self.assertFalse(manifest["preservation"]["source_padj_recomputed"])
            for figure, (_, sheet, _, _, multiplier) in zip(manifest["figures"], legacy.FIGURES):
                table = pd.read_csv(destination / figure["data_file"], sep="\t", float_precision="round_trip")
                self.assertEqual(tuple(table.columns), legacy.OUTPUT_COLUMNS)
                self.assertEqual(table.source_sheet.unique().tolist(), [sheet])
                self.assertEqual(table.log2FoldChange.tolist(), [0.4 * multiplier, -0.5 * multiplier, 0.0])
                self.assertEqual(table.source_log2FoldChange.tolist(), [0.4, -0.5, 0.0])
                self.assertEqual(table.pvalue.tolist(), [0.001, 0.01, 0.5])
                self.assertEqual(table.padj.tolist(), [0.003, 0.015, 0.5])
                self.assertEqual(figure["audit"]["selection"]["joint_deg_count"], 2)
                self.assertTrue(figure["audit"]["bh_reference"]["agrees_with_source"])
                self.assertEqual(figure["audit"]["duplicate_symbols_with_conflicting_effect_directions"], ["Shared"])
            self.assertEqual(legacy.validate_completed_import(destination), manifest)

    def test_non_bh_source_adjustments_are_reported_and_preserved(self):
        rows = source_rows()
        rows[0][7] = 0.02
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "source.xlsx"
            source.write_bytes(workbook_bytes(rows))
            destination = Path(temporary) / "imported"
            manifest = legacy.import_workbook(source, destination)
            self.assertFalse(manifest["figures"][0]["audit"]["bh_reference"]["agrees_with_source"])
            imported = pd.read_csv(destination / "figure_01_deg.tsv.gz", sep="\t")
            self.assertEqual(imported.padj.iloc[0], 0.02)

    def test_invalid_reimport_leaves_previous_complete_artifacts_unchanged(self):
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "source.xlsx"
            source.write_bytes(workbook_bytes())
            destination = Path(temporary) / "imported"
            legacy.import_workbook(source, destination)
            previous = {path.name: path.read_bytes() for path in destination.iterdir()}
            rows = source_rows()
            rows[1][6] = 0
            source.write_bytes(workbook_bytes(rows))
            with self.assertRaises(ValueError):
                legacy.import_workbook(source, destination)
            self.assertEqual(previous, {path.name: path.read_bytes() for path in destination.iterdir()})
            legacy.validate_completed_import(destination)

    def test_data_tampering_fails_closed(self):
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "source.xlsx"
            source.write_bytes(workbook_bytes())
            destination = Path(temporary) / "imported"
            legacy.import_workbook(source, destination)
            path = destination / "figure_02_deg.tsv.gz"
            path.write_bytes(path.read_bytes() + b"changed")
            with self.assertRaisesRegex(ValueError, "hash mismatch"):
                legacy.validate_completed_import(destination)

    def test_unrelated_or_synthetic_directories_are_never_overwritten(self):
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "source.xlsx"
            source.write_bytes(workbook_bytes())
            destination = Path(temporary) / "demo"
            destination.mkdir()
            sentinel = destination / "simulation_summary.json"
            sentinel.write_text('{"data_status":"Entirely synthetic illustration;"}')
            before = sentinel.read_bytes()
            with self.assertRaisesRegex(ValueError, "Refusing to overwrite"):
                legacy.import_workbook(source, destination)
            self.assertEqual(sentinel.read_bytes(), before)


if __name__ == "__main__":
    unittest.main()
