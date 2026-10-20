#!/usr/bin/env python3
"""Import the supplied legacy mouse DEG workbook without refitting an experiment.

Source P values and adjusted P values remain unchanged. The two derived rankings
are explicitly exploratory proxies, not recovered source test statistics. This
importer refuses zero probabilities rather than silently inventing a ranking
floor; upstream underflow handling must be documented before such files are used.
"""

from __future__ import annotations

import argparse
from hashlib import sha256
from io import BytesIO
import json
import os
from pathlib import Path
import re
import shutil
import tempfile

import numpy as np
import openpyxl
import pandas as pd
from scipy import stats


HEADERS = ("ENSEMBL", "SYMBOL", "ENTREZID", "baseMean", "log2FoldChange", "lfcSE", "pvalue", "padj")
FIGURES = (
    ("figure_01", "saline_VS_wt", "LD SALINE", "NON-LD control", 1),
    ("figure_02", "saline_VS_leptin", "LD LEPTIN", "LD SALINE", -1),
    ("figure_03", "leptin_VS_wt", "LD LEPTIN", "NON-LD control", 1),
)
OUTPUT_COLUMNS = (
    "gene_id", "supplied_symbol", "entrez_id", "baseMean", "log2FoldChange", "lfcSE",
    "pvalue", "padj", "source_log2FoldChange", "source_sheet",
    "primaryproxy_signed_log10_p", "sensitivityproxy_signed_z",
)
STATUS = "Imported legacy differential-expression results; no experimental refit."
MANIFEST = "legacy_provenance.json"
ENSEMBL_PATTERN = re.compile(r"ENSMUSG[0-9]+(?:\.[0-9]+)?\Z")


def _digest(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def _numeric(value: object, field: str, location: str) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"{location}: {field} must be a numeric Excel cell.")
    number = float(value)
    if not np.isfinite(number):
        raise ValueError(f"{location}: {field} must be finite.")
    return number


def read_workbook(workbook_bytes: bytes) -> dict[str, pd.DataFrame]:
    """Validate the three source sheets; never execute formulas or external links."""
    workbook = openpyxl.load_workbook(
        BytesIO(workbook_bytes), read_only=True, data_only=False, keep_links=False
    )
    try:
        expected = {figure[1] for figure in FIGURES}
        if set(workbook.sheetnames) != expected or len(workbook.sheetnames) != len(expected):
            raise ValueError(f"Expected exactly these sheets: {sorted(expected)}.")
        tables = {}
        for sheet_name in sorted(expected):
            sheet = workbook[sheet_name]
            rows = iter(sheet.iter_rows())
            header = next(rows, ())
            if tuple(cell.value for cell in header) != HEADERS:
                raise ValueError(f"{sheet_name}: expected exactly the eight documented column headers.")
            records = []
            seen = set()
            for row_number, cells in enumerate(rows, start=2):
                location = f"{sheet_name} row {row_number}"
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError(f"{location}: formulas are not accepted or executed.")
                values = [cell.value for cell in cells]
                if len(values) != len(HEADERS) or any(value is None for value in values):
                    raise ValueError(f"{location}: missing cell or incorrect column count.")
                gene, symbol, entrez, base_mean, fold_change, standard_error, pvalue, padj = values
                if not isinstance(gene, str) or not ENSEMBL_PATTERN.fullmatch(gene):
                    raise ValueError(f"{location}: expected a mouse Ensembl gene identifier.")
                if gene in seen:
                    raise ValueError(f"{location}: duplicate Ensembl gene identifier {gene}.")
                seen.add(gene)
                if not isinstance(symbol, str) or not symbol.strip() or any(c in symbol for c in "\t\r\n"):
                    raise ValueError(f"{location}: SYMBOL must be nonempty text without control characters.")
                entrez = _numeric(entrez, "ENTREZID", location)
                if entrez <= 0 or entrez != np.floor(entrez) or entrez > 2**53 - 1:
                    raise ValueError(f"{location}: ENTREZID must be an exactly representable positive integer.")
                numbers = [
                    _numeric(value, field, location)
                    for value, field in zip(
                        (base_mean, fold_change, standard_error, pvalue, padj), HEADERS[3:]
                    )
                ]
                base_mean, fold_change, standard_error, pvalue, padj = numbers
                if base_mean < 0 or standard_error <= 0:
                    raise ValueError(f"{location}: baseMean must be nonnegative and lfcSE positive.")
                if not 0 < pvalue <= 1 or not 0 < padj <= 1:
                    raise ValueError(
                        f"{location}: pvalue and padj must be in (0, 1]; zero values require "
                        "documented upstream underflow handling, not an invented ranking floor."
                    )
                records.append((gene, symbol, int(entrez), *numbers))
            if not records:
                raise ValueError(f"{sheet_name}: the source sheet has no data rows.")
            tables[sheet_name] = pd.DataFrame(records, columns=HEADERS)
        return tables
    finally:
        workbook.close()


def orient_table(source: pd.DataFrame, sheet_name: str, multiplier: int) -> pd.DataFrame:
    """Retain every Ensembl row and orient its effect to the requested contrast."""
    if multiplier not in (-1, 1):
        raise ValueError("Fold-change orientation must be +1 or -1.")
    result = source.rename(
        columns={"ENSEMBL": "gene_id", "SYMBOL": "supplied_symbol", "ENTREZID": "entrez_id"}
    ).copy()
    result["source_log2FoldChange"] = result["log2FoldChange"].copy()
    result["log2FoldChange"] = multiplier * result["log2FoldChange"]
    result["source_sheet"] = sheet_name
    direction = np.sign(result["log2FoldChange"].to_numpy())
    source_p = result["pvalue"].to_numpy()
    result["primaryproxy_signed_log10_p"] = direction * -np.log10(source_p)
    with np.errstate(under="ignore"):
        sensitivity = direction * stats.norm.isf(source_p / 2)
    if not np.all(np.isfinite(sensitivity)):
        raise ValueError("P-value underflow prevents a finite signed-z proxy; obtain upstream handling.")
    result["sensitivityproxy_signed_z"] = sensitivity
    return result.loc[:, OUTPUT_COLUMNS]


def _table_audit(source: pd.DataFrame, oriented: pd.DataFrame) -> dict:
    pvalue = source["pvalue"].to_numpy()
    padj = source["padj"].to_numpy()
    reconstructed_bh = stats.false_discovery_control(pvalue, method="bh")
    difference = np.abs(reconstructed_bh - padj)
    with np.errstate(over="ignore", divide="ignore", invalid="ignore"):
        ratio = source["log2FoldChange"].to_numpy() / source["lfcSE"].to_numpy()
        ratio_implied_log_p = np.log(2.0) + stats.norm.logsf(np.abs(ratio))
    compatible = np.isfinite(ratio_implied_log_p)
    log_error = np.abs(np.log(pvalue[compatible]) - ratio_implied_log_p[compatible]) / np.log(10)
    selected = (oriented["padj"] < 0.05) & (oriented["log2FoldChange"].abs() >= 0.2)
    duplicate_symbols = source[source["SYMBOL"].duplicated(keep=False)]
    conflicts = [
        symbol for symbol, rows in duplicate_symbols.groupby("SYMBOL", sort=True)
        if np.sign(rows["log2FoldChange"]).nunique() > 1
    ]
    return {
        "rows": len(source),
        "unique_ensembl_genes": int(source["ENSEMBL"].nunique()),
        "duplicate_symbol_extra_rows": int(source["SYMBOL"].duplicated().sum()),
        "duplicate_entrez_extra_rows": int(source["ENTREZID"].duplicated().sum()),
        "duplicate_symbols_with_conflicting_effect_directions": conflicts,
        "all_ensembl_rows_retained": bool(len(source) == len(oriented)),
        "missing_numeric_values": 0,
        "baseMean_minimum": float(source["baseMean"].min()),
        "baseMean_maximum": float(source["baseMean"].max()),
        "absolute_log2fc_median": float(source["log2FoldChange"].abs().median()),
        "absolute_log2fc_below_1e_minus5": int((source["log2FoldChange"].abs() < 1e-5).sum()),
        "bh_reference": {
            "family_size": len(source),
            "max_absolute_difference": float(difference.max()),
            "agrees_with_source": bool(np.allclose(reconstructed_bh, padj, rtol=1e-10, atol=1e-15)),
            "used_to_replace_source_padj": False,
        },
        "lfc_divided_by_se_compatibility": {
            "comparison": "Source P versus two-sided normal P from supplied log2FoldChange/lfcSE; diagnostic only.",
            "median_absolute_log10_p_difference": float(np.median(log_error)) if len(log_error) else None,
            "max_absolute_log10_p_difference": float(np.max(log_error)) if len(log_error) else None,
            "nonfinite_ratio_implied_log_p": int((~compatible).sum()),
            "is_source_test_statistic": False,
        },
        "selection": {
            "padj_below_0_05": int((oriented["padj"] < 0.05).sum()),
            "joint_deg_count": int(selected.sum()),
            "higher_expression": int((selected & (oriented["log2FoldChange"] > 0)).sum()),
            "lower_expression": int((selected & (oriented["log2FoldChange"] < 0)).sum()),
        },
    }


def _cross_contrast_audit(tables: dict[str, pd.DataFrame], oriented: dict[str, pd.DataFrame]) -> dict:
    sets = {sheet: set(table["ENSEMBL"]) for sheet, table in tables.items()}
    common = sorted(set.intersection(*sets.values()))
    union = set.union(*sets.values())
    largest_name = max(tables, key=lambda name: len(tables[name]))
    largest = tables[largest_name]
    filters = {}
    for sheet, table in tables.items():
        minimum = table["baseMean"].min()
        threshold_ids = set(largest.loc[largest["baseMean"] >= minimum, "ENSEMBL"])
        filters[sheet] = {
            "observed_baseMean_minimum": float(minimum),
            "exact_baseMean_threshold_subset_of_largest_sheet": threshold_ids == sets[sheet],
            "genes_from_largest_sheet_absent_here": len(sets[largest_name] - sets[sheet]),
            "upstream_filtering_method": "Unknown; threshold structure is an observation, not inferred pipeline metadata.",
        }
    indexed = {figure: table.set_index("gene_id").loc[common] for figure, table in oriented.items()}
    if common:
        error = indexed["figure_03"]["log2FoldChange"] - indexed["figure_01"]["log2FoldChange"] - indexed["figure_02"]["log2FoldChange"]
        means = np.stack([table["baseMean"].to_numpy() for table in indexed.values()])
        identity = {
            "definition": "(LD LEPTIN - NON-LD control) - (LD SALINE - NON-LD control) - (LD LEPTIN - LD SALINE)",
            "common_genes": len(common),
            "maximum_absolute_error": float(error.abs().max()),
            "median_absolute_error": float(error.abs().median()),
            "genes_within_1e_minus8": int((error.abs() < 1e-8).sum()),
            "maximum_baseMean_difference": float(np.ptp(means, axis=0).max()),
            "explanation": "Upstream model/shrinkage methods are unknown; do not impose additive effects or infer rescue.",
        }
    else:
        identity = {"common_genes": 0, "explanation": "No common genes; additive-effect comparison unavailable."}
    return {
        "union_genes": len(union), "common_genes": len(common),
        "largest_source_sheet": largest_name,
        "observed_filter_structure": filters,
        "contrast_identity": identity,
    }


def _write_import(
    workbook_bytes: bytes, workbook_name: str, destination: Path,
    complete_tested_universe: bool,
) -> dict:
    tables = read_workbook(workbook_bytes)
    oriented = {}
    figures = []
    for figure_id, sheet, numerator, denominator, multiplier in FIGURES:
        table = orient_table(tables[sheet], sheet, multiplier)
        oriented[figure_id] = table
        data_file = f"{figure_id}_deg.tsv.gz"
        table.to_csv(
            destination / data_file, sep="\t", index=False, float_format="%.17g",
            compression={"method": "gzip", "mtime": 0},
        )
        figures.append({
            "figure_id": figure_id, "data_file": data_file, "source_sheet": sheet,
            "numerator": numerator, "denominator": denominator,
            "fold_change_multiplier": multiplier,
            "audit": _table_audit(tables[sheet], table),
        })
    manifest = {
        "complete": True,
        "data_status": STATUS,
        "input_file": workbook_name,
        "input_sha256": sha256(workbook_bytes).hexdigest(),
        "tested_universe_confirmation": {
            "all_tested_genes_exported": bool(complete_tested_universe),
            "source": "User confirmed all tested genes were exported." if complete_tested_universe else "No confirmation supplied to importer.",
        },
        "importer_sha256": _digest(Path(__file__)),
        "source_paths": {item["figure_id"]: item["data_file"] for item in figures},
        "data_artifact_sha256": {item["data_file"]: _digest(destination / item["data_file"]) for item in figures},
        "figures": figures,
        "cross_contrast_audit": _cross_contrast_audit(tables, oriented),
        "preservation": {
            "source_pvalues_recomputed": False, "source_padj_recomputed": False,
            "experimental_refit_performed": False,
            "source_effects_retained_in_source_log2FoldChange": True,
            "ensembl_rows_dropped": 0,
            "duplicate_symbols_collapsed_in_de_tables": False,
            "source_workbook_copied_to_repository": False,
        },
        "deg_display_rule": {"padj_strictly_below": 0.05, "absolute_log2fc_at_least": 0.2},
        "ranking_proxies": {
            "primaryproxy_signed_log10_p": "sign(oriented supplied log2FoldChange) * -log10(source pvalue)",
            "sensitivityproxy_signed_z": "sign(oriented supplied log2FoldChange) * normal.isf(source pvalue / 2)",
            "status": "Derived exploratory proxies; neither is the missing original test statistic.",
            "filters": "No DEG or effect-size filter is applied to ranking inputs; every available source row is retained.",
            "zero_probability_policy": "Refuse zero P/padj; request documented upstream underflow handling. No floor is invented.",
        },
        "findings": [
            "Raw counts, sample-level expression, replicate metadata, original signed test statistics and model/shrinkage methods are unavailable.",
            "Gene universes differ by contrast and follow distinct observed expression thresholds; upstream filtering methods remain unknown.",
            "Supplied effect estimates and SEs must not be substituted for the missing original test statistic.",
            "Very small post-treatment effect estimates may reflect upstream processing; preservation does not validate biological normalization.",
            "A nonsignificant comparison is not evidence of equivalence or complete rescue.",
            "Duplicate symbol/Entrez mappings remain in DE tables; enrichment requires a separately documented mapping policy.",
        ],
        "software": {"numpy": np.__version__, "pandas": pd.__version__, "openpyxl": openpyxl.__version__},
    }
    (destination / MANIFEST).write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False, allow_nan=False) + "\n", encoding="utf-8"
    )
    return manifest


def import_workbook(
    workbook_path: Path, output_dir: Path | None = None, *, complete_tested_universe: bool = False
) -> dict:
    """Publish a validated import with rollback; do not overwrite unrelated data."""
    workbook_path = Path(workbook_path)
    workbook_bytes = workbook_path.read_bytes()
    destination = Path(output_dir or Path(__file__).resolve().parent / "legacy_results").absolute()
    if destination.is_symlink() or (destination.exists() and not destination.is_dir()):
        raise ValueError("The import destination must be a real directory, not a file or symlink.")
    if destination.exists() and any(destination.iterdir()):
        try:
            previous = json.loads((destination / MANIFEST).read_text(encoding="utf-8"))
        except (OSError, ValueError) as error:
            raise ValueError("Refusing to overwrite a nonempty directory without a legacy-import marker.") from error
        if not isinstance(previous, dict) or previous.get("data_status") != STATUS:
            raise ValueError("Refusing to overwrite data not marked as a legacy-workbook import.")
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix=f".{destination.name}-staging-", dir=destination.parent) as temporary:
        stage_root = Path(temporary)
        staged = stage_root / "results"
        if destination.exists():
            shutil.copytree(destination, staged, symlinks=True)
            if any(path.is_symlink() for path in staged.iterdir()):
                raise ValueError("Legacy-import outputs must not contain symbolic links.")
        else:
            staged.mkdir()
        manifest = _write_import(workbook_bytes, workbook_path.name, staged, complete_tested_universe)
        backup = stage_root / "previous"
        if destination.exists():
            os.replace(destination, backup)
        try:
            os.replace(staged, destination)
        except OSError:
            if backup.exists():
                os.replace(backup, destination)
            raise
    return manifest


def validate_completed_import(data_dir: Path) -> dict:
    """Fail closed on incomplete, unexpected or modified imported artifacts."""
    data_dir = Path(data_dir)
    try:
        manifest = json.loads((data_dir / MANIFEST).read_text(encoding="utf-8"))
    except (OSError, ValueError) as error:
        raise ValueError("Legacy import provenance is missing or unreadable.") from error
    if not isinstance(manifest, dict) or manifest.get("complete") is not True or manifest.get("data_status") != STATUS:
        raise ValueError("Legacy workbook import is not marked complete.")
    expected_files = {f"{item[0]}_deg.tsv.gz" for item in FIGURES}
    artifact_hashes = manifest.get("data_artifact_sha256", {})
    if not isinstance(artifact_hashes, dict) or set(artifact_hashes) != expected_files:
        raise ValueError("Legacy import artifact manifest has unexpected data files.")
    figures = manifest.get("figures", [])
    if not isinstance(figures, list) or len(figures) != len(FIGURES):
        raise ValueError("Legacy import must describe exactly three ordered figures.")
    for figure, expected in zip(figures, FIGURES):
        figure_id, sheet, numerator, denominator, multiplier = expected
        required = {
            "figure_id": figure_id, "data_file": f"{figure_id}_deg.tsv.gz",
            "source_sheet": sheet, "numerator": numerator, "denominator": denominator,
            "fold_change_multiplier": multiplier,
        }
        if not isinstance(figure, dict) or any(figure.get(key) != value for key, value in required.items()):
            raise ValueError("Legacy import contrast directions or data files have changed.")
    for filename, expected_hash in artifact_hashes.items():
        artifact = data_dir / filename
        if artifact.is_symlink() or not artifact.is_file() or _digest(artifact) != expected_hash:
            raise ValueError(f"Legacy import data hash mismatch: {filename}.")
    return manifest


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True, help="Legacy .xlsx workbook with the three documented contrasts.")
    parser.add_argument("--output-dir", type=Path, default=None, help="Defaults to analysis/legacy_results.")
    parser.add_argument(
        "--complete-tested-universe", action="store_true",
        help="Record the user's confirmation that every tested gene was exported for each contrast.",
    )
    arguments = parser.parse_args()
    manifest = import_workbook(
        arguments.workbook, arguments.output_dir, complete_tested_universe=arguments.complete_tested_universe
    )
    print(json.dumps({
        "complete": manifest["complete"], "input_file": manifest["input_file"],
        "figures": [{"figure_id": item["figure_id"], **item["audit"]["selection"]} for item in manifest["figures"]],
    }, indent=2))


if __name__ == "__main__":
    main()
